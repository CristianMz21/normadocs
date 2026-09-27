"""End-to-end tests for APA block quotes (>40 words).

Fixture: Markdown with a ``>`` blockquote over 40 words. The pipeline
(CLI -> Pandoc -> APA formatter) must emit it as its own freestanding
block with a 0.5" left indent, no first-line indent, no surrounding
quotes, and no extra blank paragraphs around it.

Requires: pandoc installed.
"""

import shutil
import textwrap
import unittest
from pathlib import Path

from docx import Document
from docx.shared import Inches
from docx.text.paragraph import Paragraph
from typer.testing import CliRunner

from normadocs.cli import app

runner = CliRunner()

PANDOC_AVAILABLE = shutil.which("pandoc") is not None

BLOCK_QUOTE_FIXTURE_MD = textwrap.dedent("""\
    ---
    title: "Prueba Cita en Bloque"
    author: "Autor Prueba"
    program: "ADSO"
    institution: "SENA"
    date: "2026-09-27"
    ---

    # Introducción

    Este es un párrafo introductorio que precede a la cita en bloque de prueba.

    > El aprendizaje automático permite personalizar el contenido educativo de forma
    > que cada estudiante avanza según sus propias necesidades y ritmo de aprendizaje
    > particular lo que a largo plazo mejora significativamente la retención y el
    > rendimiento académico medido en evaluaciones estandarizadas diversas y complejas
    > (García, 2020).

    Este es el párrafo posterior a la cita en bloque para verificar continuidad.
    """)

QUOTE_MARKERS = ('"', "\u201c", "\u201d", "\u00ab", "\u00bb")


@unittest.skipUnless(PANDOC_AVAILABLE, "Pandoc not installed — skipping block-quote E2E")
class TestBlockQuoteEndToEnd(unittest.TestCase):
    """Full pipeline: CLI -> Pandoc -> APA formatter -> DOCX block validation."""

    output_dir: Path

    @classmethod
    def setUpClass(cls) -> None:
        cls.output_dir = Path("tests/temp_block_quote_e2e")
        cls.output_dir.mkdir(parents=True, exist_ok=True)
        cls.md_file = cls.output_dir / "block_quote_long.md"
        cls.md_file.write_text(BLOCK_QUOTE_FIXTURE_MD, encoding="utf-8")
        result = runner.invoke(
            app,
            [
                str(cls.md_file),
                "--style",
                "apa",
                "--output-dir",
                str(cls.output_dir),
            ],
        )
        cls.cli_result = result
        cls.docx_path = cls.output_dir / "block_quote_long_APA.docx"
        cls.doc = Document(str(cls.docx_path)) if cls.docx_path.exists() else None

    @classmethod
    def tearDownClass(cls) -> None:
        if cls.output_dir.exists():
            shutil.rmtree(cls.output_dir)

    def _block_paragraph(self) -> Paragraph:
        self.assertIsNotNone(self.doc, "DOCX not loaded")
        assert self.doc is not None
        for para in self.doc.paragraphs:
            if "retención" in para.text and len(para.text.split()) > 40:
                return para
        self.fail("Block-quote paragraph (>40 words) not found in output DOCX")
        raise AssertionError("unreachable")

    def test_cli_exits_successfully(self) -> None:
        """CLI must exit with code 0 for the block-quote fixture."""
        self.assertEqual(
            self.cli_result.exit_code,
            0,
            f"CLI failed (exit {self.cli_result.exit_code}):\n{self.cli_result.output}"
            + (f"\n{self.cli_result.exception}" if self.cli_result.exception else ""),
        )

    def test_docx_file_created(self) -> None:
        """The .docx output file must exist on disk."""
        self.assertTrue(self.docx_path.exists())

    def test_block_quote_is_own_paragraph(self) -> None:
        """Long quote must be its own paragraph, not merged with neighbours."""
        block = self._block_paragraph()
        self.assertNotIn("párrafo introductorio", block.text)
        self.assertNotIn("verificar continuidad", block.text)

    def test_block_quote_has_half_inch_left_indent(self) -> None:
        """APA 8.27: block quote indented 0.5" from the left margin."""
        block = self._block_paragraph()
        self.assertEqual(block.paragraph_format.left_indent, Inches(0.5))
        self.assertEqual(block.paragraph_format.first_line_indent, Inches(0))

    def test_block_quote_has_no_surrounding_quotes(self) -> None:
        """APA 8.27: freestanding blocks carry no quotation marks."""
        block = self._block_paragraph()
        stripped = block.text.strip()
        self.assertFalse(stripped.startswith(QUOTE_MARKERS))
        self.assertNotIn('"', block.text)

    def test_no_extra_blanks_in_body(self) -> None:
        """No empty paragraphs may surround the block in the body."""
        self.assertIsNotNone(self.doc, "DOCX not loaded")
        assert self.doc is not None
        paras = list(self.doc.paragraphs)
        first_heading = next(
            (
                i
                for i, p in enumerate(paras)
                if p.style and p.style.name == "Heading 1" and p.text.strip()
            ),
            len(paras),
        )
        body_blanks = [i for i in range(first_heading, len(paras)) if not paras[i].text.strip()]
        self.assertEqual(body_blanks, [], f"Extra blank paragraphs at {body_blanks}")


if __name__ == "__main__":
    unittest.main()
