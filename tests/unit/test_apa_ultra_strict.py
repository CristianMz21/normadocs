"""Unit tests for APA ultra-strict formatter defaults (spec 001-apa-ultra-strict).

Covers US1: 7-field cover order, subject_code model, font profiles,
and estudiante running-head default.
"""

import unittest

from docx.oxml.ns import qn

from normadocs.config import METADATA_FIELDS
from normadocs.models import DocumentMetadata
from normadocs.preprocessor import MarkdownPreprocessor
from normadocs.standards import StandardLoader
from normadocs.standards.schema import get_default_config


class TestSubjectCodeModel(unittest.TestCase):
    """DocumentMetadata must carry the course code."""

    def test_from_dict_includes_subject_code(self) -> None:
        meta = DocumentMetadata.from_dict(
            {"title": "T", "subject": "ADSO", "subject_code": "ADSO-2026"}
        )
        self.assertEqual(meta.subject, "ADSO")
        self.assertEqual(meta.subject_code, "ADSO-2026")

    def test_metadata_fields_include_subject_code(self) -> None:
        self.assertIn("subject_code", METADATA_FIELDS)


class TestPreprocessorCoverFields(unittest.TestCase):
    """Title-page markdown must list the 7 cover fields in order."""

    def test_build_title_page_md_order(self) -> None:
        meta = DocumentMetadata(
            title="Informe",
            author="Cristian Arellano",
            program="ADSO",
            institution="SENA",
            subject="Analisis",
            subject_code="ADSO-2026",
            instructor="Docente",
            date="26 de septiembre de 2026",
        )
        md = MarkdownPreprocessor.build_title_page_md(meta)
        positions = [
            md.find("Informe"),
            md.find("Cristian Arellano"),
            md.find("ADSO"),
            md.find("SENA"),
            md.find("Analisis"),
            md.find("ADSO-2026"),
            md.find("Docente"),
            md.find("septiembre"),
        ]
        self.assertTrue(all(p >= 0 for p in positions))
        self.assertEqual(positions, sorted(positions))


class TestCoverOrder(unittest.TestCase):
    """APACoverHandler must emit program/institution/subject before date."""

    def test_content_lines_order(self) -> None:
        from docx import Document

        from normadocs.formatters.apa.apa_cover import APACoverHandler

        doc = Document()
        doc.add_paragraph("Body", style="Normal")
        handler = APACoverHandler(doc, {})
        meta = DocumentMetadata(
            title="T",
            author="A",
            program="P",
            institution="I",
            subject="S",
            subject_code="C",
            instructor="D",
            date="F",
        )
        lines = [text for text, _ in handler._build_content_lines(meta) if text]
        self.assertLess(lines.index("P"), lines.index("I"))
        self.assertLess(lines.index("I"), lines.index("S (C)"))
        self.assertLess(lines.index("S (C)"), lines.index("D"))
        self.assertLess(lines.index("D"), lines.index("F"))


class TestStyleProfiles(unittest.TestCase):
    """Font profiles and running-head defaults per spec."""

    def test_allowed_profiles_in_defaults(self) -> None:
        config = get_default_config("apa7estudiante")
        allowed = config["fonts"]["allowed"]
        names = {(a["name"], a["size"]) for a in allowed}
        self.assertIn(("Times New Roman", 12), names)
        self.assertIn(("Arial", 11), names)
        self.assertIn(("Calibri", 11), names)
        self.assertIn(("Georgia", 11), names)

    def test_estudiante_running_head_disabled(self) -> None:
        config = get_default_config("apa7estudiante")
        self.assertFalse(config["running_head"]["enabled"])

    def test_estudiante_yaml_loads(self) -> None:
        config = StandardLoader().load("apa7estudiante")
        self.assertFalse(config["running_head"]["enabled"])
        self.assertEqual(len(config["fonts"]["allowed"]), 4)


class TestOrderedLists(unittest.TestCase):
    """Ordered markdown lists must keep numbering (APA 7 sequential lists)."""

    def _doc_with_numbering(self, num_fmt: str):
        """Build a doc with one numbered paragraph backed by numbering.xml."""
        from docx import Document
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn

        doc = Document()
        numbering = doc.part.numbering_part.element
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), "900")
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), "0")
        fmt = OxmlElement("w:numFmt")
        fmt.set(qn("w:val"), num_fmt)
        lvl.append(fmt)
        abstract.append(lvl)
        numbering.append(abstract)
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), "901")
        abs_id = OxmlElement("w:abstractNumId")
        abs_id.set(qn("w:val"), "900")
        num.append(abs_id)
        numbering.append(num)

        para = doc.add_paragraph("Primero", style="Normal")
        p_pr = para._element.get_or_add_pPr()
        num_pr = OxmlElement("w:numPr")
        ilvl = OxmlElement("w:ilvl")
        ilvl.set(qn("w:val"), "0")
        num_id = OxmlElement("w:numId")
        num_id.set(qn("w:val"), "901")
        num_pr.append(ilvl)
        num_pr.append(num_id)
        p_pr.append(num_pr)
        return doc, para

    def test_ordered_list_keeps_numbering(self) -> None:
        from normadocs.formatters.apa.apa_paragraphs import APAParagraphsHandler

        doc, para = self._doc_with_numbering("decimal")
        APAParagraphsHandler(doc, {}).format_lists()
        num_pr = para._element.find(qn("w:pPr")).find(qn("w:numPr"))
        self.assertIsNotNone(num_pr)
        self.assertFalse(para.text.startswith("•"))

    def test_unordered_list_gets_bullet(self) -> None:
        from normadocs.formatters.apa.apa_paragraphs import APAParagraphsHandler

        doc, para = self._doc_with_numbering("bullet")
        APAParagraphsHandler(doc, {}).format_lists()
        num_pr = para._element.find(qn("w:pPr")).find(qn("w:numPr"))
        self.assertIsNone(num_pr)
        self.assertTrue(para.text.startswith("•"))


class TestTableBorders(unittest.TestCase):
    """APA tables must carry explicit horizontal-only borders, no verticals."""

    def test_formatter_writes_no_vertical_borders(self) -> None:
        from docx import Document

        from normadocs.formatters.apa.apa_tables import APATablesHandler

        doc = Document()
        doc.add_table(rows=2, cols=2)
        APATablesHandler(doc, {})._apply_apa_table_borders(doc.tables[0])
        from normadocs.verifier.checks.tables import TablesCheck

        tbl = doc.tables[0]._tbl
        self.assertFalse(TablesCheck._has_table_level_verticals(tbl))
        checker = TablesCheck()
        self.assertFalse(checker._has_cell_level_verticals(tbl))

    def test_verifier_flags_cell_level_verticals(self) -> None:
        from docx import Document
        from docx.oxml import OxmlElement

        from normadocs.verifier.checks.tables import TablesCheck

        doc = Document()
        doc.add_table(rows=1, cols=1)
        cell = doc.tables[0].rows[0].cells[0]
        tc_pr = cell._tc.get_or_add_tcPr()
        borders = OxmlElement("w:tcBorders")
        inside_v = OxmlElement("w:insideV")
        inside_v.set(qn("w:val"), "single")
        borders.append(inside_v)
        tc_pr.append(borders)
        checker = TablesCheck()
        self.assertTrue(checker._has_cell_level_verticals(doc.tables[0]._tbl))

    def test_table_rows_chained_with_keep_next(self) -> None:
        from docx import Document
        from docx.oxml.ns import qn as _qn

        from normadocs.formatters.apa.apa_tables import APATablesHandler

        doc = Document()
        caption = doc.add_paragraph("Tabla 1")
        title = doc.add_paragraph("Titulo")
        for run in title.runs:
            run.italic = True
        doc.add_table(rows=3, cols=2)
        APATablesHandler(doc, {})._keep_table_together(doc.tables[0])
        for para in (caption, title):
            p_pr = para._element.find(_qn("w:pPr"))
            self.assertIsNotNone(p_pr.find(_qn("w:keepNext")))
        for row in doc.tables[0].rows:
            tr_pr = row._tr.trPr
            keep = tr_pr.find(_qn("w:keepNext")) if tr_pr is not None else None
            self.assertIsNone(keep)


if __name__ == "__main__":
    unittest.main()
