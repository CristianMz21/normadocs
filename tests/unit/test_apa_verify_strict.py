"""Unit tests for APA ultra-strict verification (spec 001-apa-ultra-strict).

Covers US2/US3/US4: font profiles, headings L3, short-quote pages,
citation<->reference cross-check, table mention-before, informe order,
objetivos count, 7-field cover.
"""

import unittest
from pathlib import Path
from tempfile import TemporaryDirectory
from typing import Any
from unittest.mock import MagicMock

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from normadocs.models import DocumentMetadata
from normadocs.verifier.apa_verifier import VerificationContext
from normadocs.verifier.checks.citations import CitationsCheck
from normadocs.verifier.checks.cover_page import CoverPageCheck
from normadocs.verifier.checks.cross_refs import CrossRefsCheck
from normadocs.verifier.checks.fonts import FontsCheck
from normadocs.verifier.checks.headings import HeadingsCheck
from normadocs.verifier.checks.structure import StructureCheck
from normadocs.verifier.checks.tables import TablesCheck
from normadocs.verifier.docx_analyzer import DOCXAnalyzer, DOCXParagraphInfo


def _para(
    text: str,
    style: str | None = None,
    alignment: str = "left",
    runs: list[dict] | None = None,
) -> DOCXParagraphInfo:
    return DOCXParagraphInfo(
        text=text,
        style_name=style,
        alignment=alignment,
        first_line_indent=None,
        space_before=None,
        space_after=None,
        line_spacing=2.0,
        runs=runs if runs is not None else [],
    )


def _run_run(text: str, font: str = "Times New Roman", size_emu: int = 152400) -> dict:
    return {"text": text, "font_name": font, "font_size": size_emu}


class _Source:
    def __init__(self, paragraphs: list[DOCXParagraphInfo]) -> None:
        self._paragraphs = paragraphs

    def get_paragraphs_info(self) -> list[DOCXParagraphInfo]:
        return self._paragraphs


def _ctx(
    paragraphs: list[DOCXParagraphInfo],
    strict: bool = True,
    title: str = "Informe",
) -> VerificationContext:
    docx: Any = _Source(paragraphs)
    return VerificationContext(
        pdf=MagicMock(),
        docx=docx,
        meta=DocumentMetadata(title=title),
        strict=strict,
    )


class TestFontProfiles(unittest.TestCase):
    """Arial 11 passes strict; disallowed fonts fail."""

    def test_arial_11_passes_strict(self) -> None:
        paras = [_para("Body text.", "Normal", runs=[_run_run("Body text.", "Arial", 139700)])]
        issues = FontsCheck().run(_ctx(paras))
        self.assertEqual([i for i in issues if "profile" in i.check], [])

    def test_comic_sans_fails_strict(self) -> None:
        paras = [_para("Body text.", "Normal", runs=[_run_run("Body text.", "Comic Sans", 152400)])]
        issues = FontsCheck().run(_ctx(paras))
        self.assertTrue(any("profile_mismatch" in i.check for i in issues))


class TestHeadingsLevel3(unittest.TestCase):
    """Level-3 defects are errors in strict, warnings otherwise."""

    def _paras(self) -> list[DOCXParagraphInfo]:
        return [_para("Marco", "Heading 3", runs=[{"text": "Marco"}])]

    def test_level3_strict_is_error(self) -> None:
        issues = HeadingsCheck().run(_ctx(self._paras(), strict=True))
        self.assertTrue(any(i.severity == "error" for i in issues))

    def test_level3_lax_is_warning(self) -> None:
        issues = HeadingsCheck().run(_ctx(self._paras(), strict=False))
        self.assertTrue(all(i.severity == "warning" for i in issues))


class TestShortQuotePage(unittest.TestCase):
    """Short quotations without p./pp. fail in strict."""

    def test_short_quote_without_page_fails(self) -> None:
        paras = [_para('"Texto exacto citado" (Perez, 2020).', "Normal")]
        issues = CitationsCheck().run(_ctx(paras))
        self.assertTrue(any("page_missing" in i.check for i in issues))

    def test_short_quote_with_page_passes(self) -> None:
        paras = [_para('"Texto exacto citado" (Perez, 2020, p. 4).', "Normal")]
        issues = CitationsCheck().run(_ctx(paras))
        self.assertEqual([i for i in issues if "page_missing" in i.check], [])


class TestCrossRefs(unittest.TestCase):
    """Citations without references and vice versa both fail."""

    def _paras(self, *texts: str) -> list[DOCXParagraphInfo]:
        return [_para(t, "Normal" if i > 0 else "Heading 1") for i, t in enumerate(texts)]

    def test_citation_without_reference_fails(self) -> None:
        paras = [
            _para("Informe", "Heading 1"),
            _para("Como dice Garcia (2020, p. 1).", "Normal"),
            _para("Referencias", "Heading 1"),
            _para("Perez, A. (2019). Titulo. Editorial.", "Normal"),
        ]
        issues = CrossRefsCheck().run(_ctx(paras))
        self.assertTrue(any("reference_missing" in i.check for i in issues))

    def test_reference_without_citation_fails(self) -> None:
        paras = [
            _para("Informe", "Heading 1"),
            _para("Texto sin citas.", "Normal"),
            _para("Referencias", "Heading 1"),
            _para("Perez, A. (2019). Titulo. Editorial.", "Normal"),
        ]
        issues = CrossRefsCheck().run(_ctx(paras))
        self.assertTrue(any("uncited" in i.check for i in issues))

    def test_matching_pair_passes(self) -> None:
        paras = [
            _para("Informe", "Heading 1"),
            _para("Como dice Perez (2019, p. 1).", "Normal"),
            _para("Referencias", "Heading 1"),
            _para("Perez, A. (2019). Titulo. Editorial.", "Normal"),
        ]
        self.assertEqual(CrossRefsCheck().run(_ctx(paras)), [])

    def test_multi_author_pair_matches_first_surname(self) -> None:
        paras = [
            _para("Informe", "Heading 1"),
            _para("El proceso es sistematico (Pressman & Maxim, 2020, p. 4).", "Normal"),
            _para("Referencias", "Heading 1"),
            _para(
                "Pressman, R. S., & Maxim, B. R. (2020). Titulo. Editorial.",
                "Normal",
            ),
        ]
        self.assertEqual(CrossRefsCheck().run(_ctx(paras)), [])

    def test_narrative_spanish_conjunction_matches(self) -> None:
        paras = [
            _para("Informe", "Heading 1"),
            _para("Garcia y Perez (2024) demostraron el efecto.", "Normal"),
            _para("Referencias", "Heading 1"),
            _para("Garcia, J., & Perez, M. (2024). Titulo. Revista.", "Normal"),
        ]
        self.assertEqual(CrossRefsCheck().run(_ctx(paras)), [])


class TestInformeStructure(unittest.TestCase):
    """Introduction optional; informe order and objetivos enforced."""

    def _base(self, *heads: str) -> list[DOCXParagraphInfo]:
        paras = [
            _para("Informe", None),
            _para("Autor", None),
            _para("2026", None),
            _para("Informe", "Heading 1"),
        ]
        for h in heads:
            paras.append(_para(h, "Heading 1"))
            paras.append(_para("Contenido.", None))
        return paras

    def test_no_introduction_passes_required(self) -> None:
        paras = self._base("Desarrollo", "Conclusiones", "Referencias")
        issues = StructureCheck().run(_ctx(paras))
        self.assertEqual([i for i in issues if i.check.endswith("_present")], [])

    def test_informe_disorder_fails(self) -> None:
        paras = self._base("Discusion", "Resultados", "Conclusiones", "Referencias")
        issues = StructureCheck().run(_ctx(paras))
        self.assertTrue(any("informe_order" in i.check for i in issues))

    def test_two_specific_objectives_fail(self) -> None:
        paras = self._base(
            "Planteamiento del problema",
            "Justificacion",
            "Objetivos",
            "Objetivo general",
            "Objetivos especificos",
            "Resultados",
            "Discusion",
            "Conclusiones",
            "Referencias",
        )
        # Insert two numbered objectives after "Objetivos especificos".
        idx = next(i for i, p in enumerate(paras) if p.text == "Objetivos especificos")
        paras[idx + 2 : idx + 2] = [_para("1. Primero.", None), _para("2. Segundo.", None)]
        issues = StructureCheck().run(_ctx(paras))
        self.assertTrue(any("informe_objetivos_count" in i.check for i in issues))

    def test_word_numbered_objectives_count(self) -> None:
        """Items with Word numbering (no visible digits) count as objectives."""
        paras = self._base(
            "Objetivos",
            "Objetivo general",
            "Objetivos especificos",
            "Resultados",
            "Discusion",
            "Conclusiones",
            "Referencias",
        )
        idx = next(i for i, p in enumerate(paras) if p.text == "Objetivos especificos")
        items = [
            DOCXParagraphInfo(
                text=t,
                style_name="Compact",
                alignment="left",
                first_line_indent=None,
                space_before=None,
                space_after=None,
                line_spacing=2.0,
                runs=[],
                is_list_item=True,
            )
            for t in ("Primero.", "Segundo.", "Tercero.")
        ]
        paras[idx + 2 : idx + 2] = items
        issues = StructureCheck().run(_ctx(paras))
        self.assertEqual([i for i in issues if "informe_objetivos" in i.check], [])


class TestTablesMentionedBefore(unittest.TestCase):
    """Tables need an in-text mention before the caption."""

    def _real_ctx(self, with_mention: bool) -> tuple[VerificationContext, TemporaryDirectory]:
        tmp = TemporaryDirectory()
        path = Path(tmp.name) / "t.docx"
        doc = Document()
        if with_mention:
            doc.add_paragraph("Como muestra la Tabla 1, los datos crecen.")
        doc.add_paragraph("Tabla 1")
        doc.add_table(rows=1, cols=1)
        doc.save(str(path))
        analyzer = DOCXAnalyzer(path)
        ctx = VerificationContext(
            pdf=MagicMock(), docx=analyzer, meta=DocumentMetadata(title="T"), strict=True
        )
        return ctx, tmp

    def test_unmentioned_table_fails(self) -> None:
        ctx, tmp = self._real_ctx(with_mention=False)
        try:
            issues = TablesCheck().run(ctx)
            self.assertTrue(any("not_cited_before" in i.check for i in issues))
        finally:
            tmp.cleanup()

    def test_mentioned_table_passes(self) -> None:
        ctx, tmp = self._real_ctx(with_mention=True)
        try:
            issues = TablesCheck().run(ctx)
            self.assertEqual([i for i in issues if "not_cited_before" in i.check], [])
        finally:
            tmp.cleanup()


class TestCoverSevenFields(unittest.TestCase):
    """Cover verifier requires program, subject, instructor in strict."""

    def _real_ctx(self, meta: DocumentMetadata) -> tuple[VerificationContext, TemporaryDirectory]:
        tmp = TemporaryDirectory()
        path = Path(tmp.name) / "c.docx"
        doc = Document()
        for text in [
            meta.title,
            meta.author or "",
            meta.program or "",
            meta.institution or "",
            meta.subject or "",
            meta.instructor or "",
            meta.date or "",
        ]:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(text)
            run.font.size = Pt(12)
        doc.add_paragraph(meta.title, style="Heading 1")
        doc.save(str(path))
        analyzer = DOCXAnalyzer(path)
        ctx = VerificationContext(pdf=MagicMock(), docx=analyzer, meta=meta, strict=True)
        return ctx, tmp

    def test_full_cover_passes(self) -> None:
        meta = DocumentMetadata(
            title="Informe",
            author="Autor",
            program="Programa",
            institution="Institucion",
            subject="Asignatura",
            subject_code="C-1",
            instructor="Docente",
            date="2026",
        )
        ctx, tmp = self._real_ctx(meta)
        try:
            issues = CoverPageCheck().run(ctx)
            missing = [i for i in issues if i.check.endswith("_present")]
            self.assertEqual(missing, [])
        finally:
            tmp.cleanup()

    def test_missing_subject_fails(self) -> None:
        meta = DocumentMetadata(
            title="Informe",
            author="Autor",
            program="Programa",
            institution="Institucion",
            subject="Asignatura",
            instructor="Docente",
            date="2026",
        )
        ctx, tmp = self._real_ctx(meta)
        # Blank the subject paragraph to simulate a missing course line.
        try:
            for p in ctx.docx.paragraphs:
                if p.text == "Asignatura":
                    p.text = ""
            issues = CoverPageCheck().run(ctx)
            self.assertTrue(any("subject_present" in i.check for i in issues))
        finally:
            tmp.cleanup()


class TestTitleCase(unittest.TestCase):
    """Cover titles should use Title Case (warning, never rewritten)."""

    def _cover_ctx(self, title: str) -> tuple[VerificationContext, TemporaryDirectory]:
        tmp = TemporaryDirectory()
        path = Path(tmp.name) / "t.docx"
        doc = Document()
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(title)
        run.bold = True
        run.font.size = Pt(12)
        doc.add_paragraph(title, style="Heading 1")
        doc.save(str(path))
        analyzer = DOCXAnalyzer(path)
        meta = DocumentMetadata(title=title)
        return (
            VerificationContext(pdf=MagicMock(), docx=analyzer, meta=meta, strict=False),
            tmp,
        )

    def test_sentence_case_warns(self) -> None:
        ctx, tmp = self._cover_ctx("Análisis de la implementación de sistemas")
        try:
            issues = CoverPageCheck().run(ctx)
            self.assertTrue(any("title_case" in i.check for i in issues))
        finally:
            tmp.cleanup()

    def test_title_case_with_minor_words_passes(self) -> None:
        ctx, tmp = self._cover_ctx("Análisis de la Implementación de Sistemas")
        try:
            issues = CoverPageCheck().run(ctx)
            self.assertEqual([i for i in issues if "title_case" in i.check], [])
        finally:
            tmp.cleanup()


class TestBookTitleItalic(unittest.TestCase):
    """Book entries without italics are auto-italicized and flagged."""

    def test_formatter_italicizes_book_title(self) -> None:
        from docx import Document as DocxDocument

        from normadocs.formatters.apa.apa_citations import APACitationsHandler

        doc = DocxDocument()
        doc.add_paragraph("Referencias", style="Heading 1")
        entry = doc.add_paragraph("Pressman, R. S. (2020). Software engineering. McGraw-Hill.")
        APACitationsHandler(doc, {}).format_references()
        self.assertTrue(any(r.italic for r in entry.runs if r.text.strip()))

    def test_verifier_flags_plain_book_entry(self) -> None:
        from normadocs.verifier.checks.references import ReferencesCheck

        paras = [
            _para("Referencias", "Heading 1"),
            _para("Pressman, R. S. (2020). Software engineering. McGraw-Hill.", None),
        ]
        issues = ReferencesCheck().run(_ctx(paras, strict=False))
        self.assertTrue(any("book_title_italic" in i.check for i in issues))


if __name__ == "__main__":
    unittest.main()
