"""Unit tests for APA abstract heading page break (R1).

The Resumen/Abstract heading must start on a new page, using the same
``_set_page_break_before`` pattern as the References heading.
"""

import unittest

from docx import Document

from normadocs.formatters.apa.apa_paragraphs import APAParagraphsHandler


class TestAbstractHeadingPageBreak(unittest.TestCase):
    """Abstract heading forces a page break like References does."""

    def test_abstract_heading_sets_page_break_before(self) -> None:
        """Resumen heading gets page_break_before via the shared setter."""
        doc = Document()
        doc.add_paragraph("Resumen", style="Heading 1")
        doc.add_paragraph("Contenido del resumen.", style="Normal")

        APAParagraphsHandler(doc, {}).process()

        heading = doc.paragraphs[0]
        self.assertTrue(heading.paragraph_format.page_break_before)

    def test_abstract_heading_does_not_stack_on_explicit_break(self) -> None:
        """An explicit preprocessor break before Resumen is not duplicated."""
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn

        doc = Document()
        heading = doc.add_paragraph("Resumen", style="Heading 1")
        break_paragraph = OxmlElement("w:p")
        break_run = OxmlElement("w:r")
        explicit_break = OxmlElement("w:br")
        explicit_break.set(qn("w:type"), "page")
        break_run.append(explicit_break)
        break_paragraph.append(break_run)
        heading._element.addprevious(break_paragraph)
        handler = APAParagraphsHandler(doc, {})

        from normadocs.formatters.apa.apa_paragraphs import ParagraphState

        handler._handle_abstract_heading(heading, ParagraphState())

        self.assertTrue(handler._has_page_break_before(heading))
        self.assertFalse(heading.paragraph_format.page_break_before)


if __name__ == "__main__":
    unittest.main()
