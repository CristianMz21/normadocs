"""Tests for utils.docx_helpers page-break scan."""

import unittest

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from normadocs.utils.docx_helpers import has_page_break_before


def _add_break_before(paragraph):
    """Insert an explicit page-break paragraph before the given one."""
    br_para = OxmlElement("w:p")
    br_run = OxmlElement("w:r")
    br = OxmlElement("w:br")
    br.set(qn("w:type"), "page")
    br_run.append(br)
    br_para.append(br_run)
    paragraph._element.addprevious(br_para)
    return br_para


def _add_bookmark_before(paragraph, name="toc-anchor"):
    """Insert pandoc-style bookmark anchors immediately before a heading."""
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), "7")
    start.set(qn("w:name"), name)
    paragraph._element.addprevious(start)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), "7")
    paragraph._element.addprevious(end)


class TestHasPageBreakBefore(unittest.TestCase):
    """Backward scan for explicit page breaks."""

    def test_break_directly_before_heading(self):
        """An explicit break paragraph counts."""
        doc = Document()
        heading = doc.add_paragraph("Conclusiones", style="Heading 1")
        _add_break_before(heading)

        self.assertTrue(has_page_break_before(heading))

    def test_bookmark_between_break_and_heading(self):
        """Pandoc bookmark anchors must not hide an explicit break."""
        doc = Document()
        heading = doc.add_paragraph("Conclusiones", style="Heading 1")
        _add_break_before(heading)
        _add_bookmark_before(heading)

        self.assertTrue(has_page_break_before(heading))

    def test_content_before_bookmark_means_no_break(self):
        """A bookmark after body text is not a page break."""
        doc = Document()
        doc.add_paragraph("Body text.", style="Normal")
        heading = doc.add_paragraph("Conclusiones", style="Heading 1")
        _add_bookmark_before(heading)

        self.assertFalse(has_page_break_before(heading))

    def test_empty_paragraphs_still_skipped(self):
        """Empty paragraphs between break and heading are transparent."""
        doc = Document()
        heading = doc.add_paragraph("Conclusiones", style="Heading 1")
        _add_break_before(heading)
        doc.paragraphs[0]._element.addnext(OxmlElement("w:p"))

        self.assertTrue(has_page_break_before(heading))

    def test_table_blocks_backward_scan(self):
        """A table is real layout: an older break does not count."""
        doc = Document()
        heading = doc.add_paragraph("Conclusiones", style="Heading 1")
        _add_break_before(heading)
        table = doc.add_table(rows=1, cols=1)
        heading._element.addprevious(table._element)

        self.assertFalse(has_page_break_before(heading))

    def test_first_paragraph_has_no_break(self):
        """Start of body never counts as a preceding break."""
        doc = Document()
        heading = doc.add_paragraph("Conclusiones", style="Heading 1")

        self.assertFalse(has_page_break_before(heading))


if __name__ == "__main__":
    unittest.main()
