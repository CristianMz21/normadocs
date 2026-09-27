"""Unit tests for BlankLinesCheck - APA 7th Edition blank-paragraph verification."""

import unittest
from pathlib import Path
from tempfile import TemporaryDirectory

from docx import Document

from normadocs.models import DocumentMetadata
from normadocs.verifier.apa_verifier import APAVerifier, VerificationContext
from normadocs.verifier.checks.blank_lines import BlankLinesCheck


class TestBlankLinesCheck(unittest.TestCase):
    """Positive, negative, and exception cases for blank-line detection."""

    @classmethod
    def setUpClass(cls) -> None:
        cls.temp_dir = TemporaryDirectory()
        cls.temp_path = Path(cls.temp_dir.name)

    @classmethod
    def tearDownClass(cls) -> None:
        cls.temp_dir.cleanup()

    def _run_check(self, doc: Document, strict: bool = True) -> list:
        path = self.temp_path / "blank_lines.docx"
        doc.save(str(path))
        pdf_path = self.temp_path / "output.pdf"
        pdf_path.touch()
        meta = DocumentMetadata(title="Test Document")
        verifier = APAVerifier(pdf_path=pdf_path, docx_path=path, meta=meta)
        ctx = VerificationContext(
            pdf=verifier.pdf,
            docx=verifier.docx,
            meta=meta,
            strict=strict,
        )
        try:
            return BlankLinesCheck().run(ctx)
        finally:
            verifier.close()

    def test_clean_document_reports_no_issues(self) -> None:
        """A document without empty body paragraphs passes cleanly."""
        doc = Document()
        doc.add_paragraph("Introducción", style="Heading 1")
        doc.add_paragraph("Primer párrafo del cuerpo.", style="Normal")
        doc.add_paragraph("Segundo párrafo del cuerpo.", style="Normal")

        issues = self._run_check(doc)

        self.assertEqual(issues, [])

    def test_blank_between_title_and_paragraph_reports_one_issue(self) -> None:
        """An empty paragraph between title and body is a single issue."""
        doc = Document()
        doc.add_paragraph("Introducción", style="Heading 1")
        doc.add_paragraph("", style="Normal")
        doc.add_paragraph("Párrafo posterior al blanco.", style="Normal")

        issues = self._run_check(doc)

        self.assertEqual(len(issues), 1)
        self.assertIn("blank_line", issues[0].check)
        self.assertEqual(issues[0].severity, "error")

    def test_blank_reports_warning_in_normal_mode(self) -> None:
        """Same blank is a warning when strict mode is off."""
        doc = Document()
        doc.add_paragraph("Introducción", style="Heading 1")
        doc.add_paragraph("", style="Normal")
        doc.add_paragraph("Párrafo posterior al blanco.", style="Normal")

        issues = self._run_check(doc, strict=False)

        self.assertEqual(len(issues), 1)
        self.assertEqual(issues[0].severity, "warning")

    def test_section_break_carrier_is_excluded(self) -> None:
        """Empty with page_break_before before Referencias is intentional."""
        doc = Document()
        doc.add_paragraph("Introducción", style="Heading 1")
        doc.add_paragraph("Cuerpo.", style="Normal")
        carrier = doc.add_paragraph("", style="Normal")
        carrier.paragraph_format.page_break_before = True
        doc.add_paragraph("Referencias", style="Heading 1")

        issues = self._run_check(doc)

        self.assertEqual(issues, [])

    def test_cover_blanks_are_excluded(self) -> None:
        """Empty paragraphs before the first Heading 1 belong to the cover."""
        doc = Document()
        doc.add_paragraph("Título del trabajo", style="Title")
        doc.add_paragraph("", style="Normal")
        doc.add_paragraph("Autor", style="Normal")
        doc.add_paragraph("Introducción", style="Heading 1")
        doc.add_paragraph("Cuerpo.", style="Normal")

        issues = self._run_check(doc)

        self.assertEqual(issues, [])

    def test_table_adjacent_blank_is_excluded(self) -> None:
        """Empty paragraph immediately before a table is table-associated."""
        doc = Document()
        doc.add_paragraph("Introducción", style="Heading 1")
        doc.add_paragraph("Cuerpo previo.", style="Normal")
        doc.add_paragraph("", style="Normal")
        doc.add_table(rows=2, cols=2)
        doc.add_paragraph("Cuerpo posterior.", style="Normal")

        issues = self._run_check(doc)

        self.assertEqual(issues, [])


if __name__ == "__main__":
    unittest.main()
